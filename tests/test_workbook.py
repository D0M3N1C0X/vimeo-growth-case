"""The workbook reproduces the Python pipeline, formula by formula."""
import pytest

formulas = pytest.importorskip("formulas")

import build_workbook  # noqa: E402
import check_workbook  # noqa: E402
from deterministic import same_content  # noqa: E402


@pytest.fixture(scope="module")
def built(a, bc, ex, tmp_path_factory):
    path = tmp_path_factory.mktemp("w") / "case.xlsx"
    return build_workbook.build(a, bc, ex, path), path


@pytest.fixture(scope="module")
def values(built):
    return check_workbook.values_from_formulas(built[1])


def test_every_formula_matches_python(built, values, capsys):
    w, _ = built
    assert check_workbook.check(values, w.wb.sheetnames, len(w.checks)) == 0, capsys.readouterr().out


def test_a_wrong_value_is_caught(built, tmp_path):
    from openpyxl import load_workbook
    w, path = built
    wb = load_workbook(path)
    ws = wb["Reconciliation"]
    big = lambda r: isinstance(ws[f"D{r}"].value, float) and abs(ws[f"D{r}"].value) > 1
    row = next(r for r in range(7, 7 + len(w.checks)) if big(r))
    ws[f"D{row}"].value *= 1.01
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    assert check_workbook.check(check_workbook.values_from_formulas(broken), wb.sheetnames, len(w.checks)) == 1


def test_changing_an_input_moves_the_case(built, tmp_path):
    from openpyxl import load_workbook
    _, path = built
    wb = load_workbook(path)
    wb["Inputs"]["C6"].value = 0.02           # uplift doubled
    moved = tmp_path / "moved.xlsx"
    wb.save(moved)
    base = check_workbook.values_from_formulas(path)[("BUSINESS CASE", "B30")]
    assert check_workbook.values_from_formulas(moved)[("BUSINESS CASE", "B30")] > base


def test_reproducible(a, bc, ex, tmp_path):
    p, q = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    build_workbook.build(a, bc, ex, p)
    build_workbook.build(a, bc, ex, q)
    assert same_content(p.read_bytes(), q.read_bytes())
