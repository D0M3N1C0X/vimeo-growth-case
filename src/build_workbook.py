"""
Builds deliverables/vimeo_upgrade_case.xlsx: the business case and the test, as live formulas.

Change an assumption on Inputs and the three-year case, the sensitivity table, the break-even uplift,
the sample sizes and the test readout recalculate. The filings data sit on their own sheets, and the
revenue bridge is computed from them. The Reconciliation sheet checks every figure against Python.

Conventions as in the other repositories: blue = input, black = formula, green = link to another
sheet, yellow = key assumption, Arial, no dynamic arrays. NORMSINV rather than NORM.S.INV, so the
file works in every Excel version, LibreOffice and Google Sheets.
"""
import math
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.hyperlink import Hyperlink

import analysis
import business_case
import config as C
import experiment
from deterministic import normalise

OUTPUT = C.DELIVERABLES / "vimeo_upgrade_case.xlsx"
REPO = "https://github.com/D0M3N1C0X/vimeo-growth-case"

FONT = "Arial"
BLUE, GREEN, INK, MUTED, WHITE = "0000FF", "008000", "1B2430", "5F6B7A", "FFFFFF"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
RED_FILL = PatternFill("solid", fgColor="F5C6C2")
GREEN_FILL = PatternFill("solid", fgColor="D5ECDC")
USD = '"$"#,##0;-"$"#,##0;"-"'
USDM = '"$"#,##0.0,,"M";-"$"#,##0.0,,"M";"-"'
NUM, NUM1, NUM2, PCT, PCT1, PCT2, PCT3 = "#,##0", "#,##0.0", "#,##0.00", "0%", "0.0%", "0.00%", "0.000%"


def font(color=INK, bold=False, italic=False, size=10):
    return Font(name=FONT, color=color, bold=bold, italic=italic, size=size)


def put(ws, ref, value, *, color=INK, bold=False, italic=False, size=10, fmt=None, fill=None, wrap=False):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold, italic, size)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def header(ws, row, labels, start=1, height=32):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=label)
        c.font = font(WHITE, bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None):
    put(ws, "A1", text, bold=True, size=14)
    if subtitle:
        put(ws, "A2", subtitle, color=MUTED, italic=True)


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def plain(x):
    if hasattr(x, "item"):
        x = x.item()
    if isinstance(x, float) and math.isnan(x):
        return None
    return x


class Workbook_:
    def __init__(self, a: dict, bc: dict, ex: dict):
        self.a, self.bc, self.ex = a, bc, ex
        self.wb = Workbook()
        self.R = {}
        self.checks = []

    def check(self, area, item, value, ref):
        self.checks.append((area, item, plain(value), ref))

    def build(self, path: Path):
        names = ["Cover", "Summary", "Business case", "Sensitivity", "Test design", "Test readout", "Value of test", "Bridge",
                 "Segments", "P&L", "Inputs", "Reconciliation"]
        self.wb.active.title = names[0]
        for n in names[1:]:
            self.wb.create_sheet(n)
        self.inputs(self.wb["Inputs"])
        self.segments(self.wb["Segments"])
        self.pnl(self.wb["P&L"])
        self.bridge(self.wb["Bridge"])
        self.case(self.wb["Business case"])
        self.sensitivity(self.wb["Sensitivity"])
        self.design(self.wb["Test design"])
        self.readout(self.wb["Test readout"])
        self.voi(self.wb["Value of test"])
        self.summary(self.wb["Summary"])
        self.reconciliation(self.wb["Reconciliation"])
        self.cover(self.wb["Cover"], names)
        for ws in self.wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.properties.creator = "Domenico Perroni"
        self.wb.properties.title = "Vimeo: the Self-Serve to Enterprise upgrade case"
        self.wb.properties.created = self.wb.properties.modified = datetime(2026, 1, 1)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        normalise(path)

    # ---- Inputs --------------------------------------------------------------------------------
    def inputs(self, ws):
        title(ws, "Inputs", "Blue = input. Yellow = the assumptions the result depends on most. Sources in docs/assumptions.md.")
        widths(ws, {"A": 3, "B": 50, "C": 16, "D": 80})
        header(ws, 4, ["", "Assumption", "Value", "Why"], height=20)
        b = C.BET
        rows = [
            ("eligible_share", "Share of Self-Serve accounts with team-level signals", b["eligible_share"], PCT1, YELLOW,
             "Seats, single sign-on interest, bandwidth, company domains: the accounts sales would call."),
            ("uplift", "Extra upgrades per eligible account per year, caused by outreach", b["uplift"], PCT2, YELLOW,
             "The unknown the test measures. 1% a year is the plan, not a forecast."),
            ("entry_acv_share", "First-year contract value, share of average Enterprise ARPU", b["entry_acv_share"], PCT, YELLOW,
             "Upgrades from Self-Serve start smaller than the average Enterprise account."),
            ("retention", "Net revenue retention of upgraded accounts after year one", b["retention"], PCT, None, ""),
            ("cannibalised", "Upgrades stop paying their Self-Serve plan (1 = yes)", int(b["cannibalised_self_serve"]), "0", None, ""),
            ("ae_cost", "Fully loaded cost of one account executive per year", b["ae_cost"], USD, None, ""),
            ("upgrades_per_ae", "Upgrades one account executive closes per year", b["upgrades_per_ae"], NUM, None, ""),
            ("planned_uplift", "Uplift the team is hired for", b["planned_uplift"], PCT2, None,
             "The team is hired before the real uplift is known: a fixed cost."),
            ("outreach_cost", "Outreach cost per eligible account per year", b["outreach_cost"], USD, YELLOW,
             "Sales development time, in-product prompts, tooling."),
            ("gross_margin", "Gross margin on the new revenue", b["gross_margin"], PCT, None, "Vimeo's gross margin was 78% in 2024."),
            ("build_cost", "One-off build of scoring and routing", b["build_cost"], USD, None, ""),
            ("discount_rate", "Discount rate", b["discount_rate"], PCT, None, ""),
            ("alpha", "Test: significance level", C.TEST["alpha"], PCT, None, ""),
            ("power", "Test: power", C.TEST["power"], PCT, None, ""),
            ("baseline_rate", "Test: organic upgrades per eligible account in 90 days", C.TEST["baseline_upgrade_rate"], PCT2, YELLOW,
             "Not public: to be read from Vimeo's own data before the test."),
            ("guardrail_churn", "Test: Self-Serve cancellations in 90 days", C.TEST["guardrail_churn"], PCT1, None, ""),
            ("guardrail_margin", "Test: largest rise in cancellations accepted, per 90 days", C.TEST["guardrail_margin"], PCT2, None, ""),
        ]
        for r, (key, label, value, fmt, fill, note) in enumerate(rows, start=5):
            put(ws, f"B{r}", label)
            put(ws, f"C{r}", value, color=BLUE, fmt=fmt, fill=fill)
            put(ws, f"D{r}", note, color=MUTED, italic=True)
            self.R[key] = f"Inputs!$C${r}"

    # ---- Filings data -------------------------------------------------------------------------
    def segments(self, ws):
        m = analysis.metrics()
        title(ws, "Segments: operating metrics from Vimeo's filings",
              "Transcribed from the MD&A of each 10-K and 10-Q; accession numbers in data/operating_metrics.csv. Two bases: do not join them.")
        widths(ws, {"A": 34, "B": 8, "C": 10, "D": 24, "E": 14, "F": 14, "G": 12, "H": 14, "I": 14, "J": 14, "K": 12})
        header(ws, 4, ["Key", "Basis", "Period", "Segment", "Subscribers (k)", "Average subscribers (k)", "ARPU ($)",
                       "Bookings ($k)", "Revenue ($k)", "ARPU x average ($k)", "Gap"])
        piv = m.pivot_table(index=["basis", "period", "segment"], columns="metric", values="value", aggfunc="first")
        order = {"Self-Serve & Add-Ons": 0, "Self-Serve": 0, "Vimeo Enterprise": 1, "Other": 2, "OTT": 2, "Add-Ons": 3}
        idx = sorted(piv.index, key=lambda k: (k[0], k[1], order.get(k[2], 9)))
        r = 5
        first = r
        for basis, period, seg in idx:
            row = piv.loc[(basis, period, seg)]
            put(ws, f"A{r}", f'=B{r}&"|"&C{r}&"|"&D{r}')
            put(ws, f"B{r}", basis, color=BLUE)
            put(ws, f"C{r}", period, color=BLUE)
            put(ws, f"D{r}", seg, color=BLUE)
            for letter, met, fmt in (("E", "subscribers_k", NUM1), ("F", "avg_subscribers_k", NUM1), ("G", "arpu_usd", NUM),
                                     ("H", "bookings_k", NUM), ("I", "revenue_k", NUM)):
                v = row.get(met)
                if v is not None and not (isinstance(v, float) and math.isnan(v)):
                    put(ws, f"{letter}{r}", float(v), color=BLUE, fmt=fmt)
            if not math.isnan(row.get("arpu_usd", float("nan"))):
                d, y = analysis.days(period)
                put(ws, f"J{r}", f"=G{r}*F{r}*{d}/{y}", fmt=NUM)
                put(ws, f"K{r}", f"=J{r}/I{r}-1", fmt=PCT2)
                if r % 3 == 0:
                    self.check("Segments", f"{period} {seg} ARPU x average", analysis.implied_revenue_k(
                        float(row["arpu_usd"]), float(row["avg_subscribers_k"]), period), f"Segments!J{r}")
            r += 1
        last = r - 1
        S = lambda letter: f"Segments!${letter}${first}:${letter}${last}"
        self.R.update({"sg_key": S("A"), "sg_avg": S("F"), "sg_arpu": S("G"), "sg_subs": S("E"), "sg_rev": S("I"),
                       "sg_book": S("H")})
        put(ws, f"D{last + 2}", "Largest gap", bold=True)
        put(ws, f"K{last + 2}", f"=MAX(MAX(K{first}:K{last}),-MIN(K{first}:K{last}))", fmt=PCT2, bold=True)
        self.check("Segments", "Largest ARPU x average gap", float(self.a["transcription"]["gap"].abs().max()), f"Segments!K{last + 2}")
        ws.freeze_panes = "E5"

    def lk(self, key, metric_range, basis, period, seg):
        return f'INDEX({metric_range},MATCH("{basis}|{period}|{seg}",{self.R["sg_key"]},0))'

    def pnl(self, ws):
        f = analysis.financials()
        p = self.a["pnl"]
        title(ws, "P&L from XBRL facts", "US$ millions. Blue figures from data/sec_financials.csv; ratios are formulas.")
        cols = list(p.index)
        widths(ws, {"A": 30, **{col(2 + i): 11 for i in range(len(cols))}})
        header(ws, 4, ["", *cols], height=20)
        lines = [("Revenue", C.REVENUE), ("Cost of revenue", "CostOfGoodsAndServicesSold"),
                 ("Sales and marketing", "SellingAndMarketingExpense"), ("Research and development", "ResearchAndDevelopmentExpense"),
                 ("General and administrative", "GeneralAndAdministrativeExpense"), ("Operating income", "OperatingIncomeLoss"),
                 ("Advertising", "AdvertisingExpense")]
        for r, (label, tag) in enumerate(lines, start=5):
            put(ws, f"A{r}", label, bold=label in ("Revenue", "Operating income"))
            for i, per in enumerate(cols):
                frame = f"CY{per}"
                if frame in f.index and tag in f.columns and not math.isnan(f.loc[frame, tag]):
                    put(ws, f"{col(2 + i)}{r}", float(f.loc[frame, tag]) / 1e6, color=BLUE, fmt=NUM1)
        ratios = [("Gross margin", lambda c: f"=1-{c}6/{c}5"), ("Sales and marketing, % revenue", lambda c: f"={c}7/{c}5"),
                  ("Research and development, % revenue", lambda c: f"={c}8/{c}5"),
                  ("General and administrative, % revenue", lambda c: f"={c}9/{c}5"),
                  ("Operating margin", lambda c: f"={c}10/{c}5"), ("Advertising, % revenue", lambda c: f'=IF({c}11="","",{c}11/{c}5)')]
        keys = ["gross_margin", "sales_marketing_pct", "research_development_pct", "general_admin_pct", "operating_income_pct", "advertising_pct"]
        for k, (label, fn) in enumerate(ratios):
            r = 13 + k
            put(ws, f"A{r}", label)
            for i, per in enumerate(cols):
                c = col(2 + i)
                put(ws, f"{c}{r}", fn(c), fmt=PCT1)
                v = p.loc[per, keys[k]]
                if not math.isnan(v):
                    self.check("P&L", f"{per} {label}", float(v), f"'P&L'!{c}{r}")
        self.R["gross_margin_2024"] = "'P&L'!$E$13"

    def bridge(self, ws):
        b = self.a["bridge"]
        title(ws, "Revenue bridge: volume and price",
              "Change in revenue = change in average subscribers at the mid ARPU (volume) + change in ARPU at the mid subscriber count (price), "
              "with ARPU brought back to each period's days; the rest is the rounding of published ARPU. US$ millions.")
        widths(ws, {"A": 10, "B": 10, "C": 24, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12})
        header(ws, 4, ["From", "To", "Segment", "Revenue from", "Revenue to", "Volume", "Price", "Rounding", "Change"])
        R_ = self.R
        for r, (_, t) in enumerate(b.iterrows(), start=5):
            frm, to, seg = t["from"], t["to"], t["segment"]
            basis = "2022" if len(to) == 4 else "2025"
            put(ws, f"A{r}", frm, color=BLUE)
            put(ws, f"B{r}", to, color=BLUE)
            put(ws, f"C{r}", seg, color=BLUE)
            g = lambda rng, per: self.lk(None, rng, basis, per, seg)
            d0, y0 = analysis.days(frm)
            d1, y1 = analysis.days(to)
            u0 = f"{g(R_['sg_arpu'], frm)}*{d0}/{y0}"         # revenue per average subscriber in the period
            u1 = f"{g(R_['sg_arpu'], to)}*{d1}/{y1}"
            put(ws, f"D{r}", f"={g(R_['sg_rev'], frm)}/1000", color=GREEN, fmt=NUM2)
            put(ws, f"E{r}", f"={g(R_['sg_rev'], to)}/1000", color=GREEN, fmt=NUM2)
            put(ws, f"F{r}", f"=({g(R_['sg_avg'], to)}-{g(R_['sg_avg'], frm)})*({u0}+{u1})/2/1000", fmt=NUM2)
            put(ws, f"G{r}", f"=({u1}-{u0})*({g(R_['sg_avg'], frm)}+{g(R_['sg_avg'], to)})/2/1000", fmt=NUM2)
            put(ws, f"H{r}", f"=E{r}-D{r}-F{r}-G{r}", fmt=NUM2)
            put(ws, f"I{r}", f"=E{r}-D{r}", fmt=NUM2, bold=True)
            for letter, name in (("F", "volume"), ("G", "price"), ("I", "change")):
                self.check("Bridge", f"{frm}-{to} {seg} {name}", float(t[name]), f"Bridge!{letter}{r}")
        ws.freeze_panes = "D5"

    # ---- The case ------------------------------------------------------------------------------
    def case(self, ws):
        R_, base, m = self.R, self.a["baseline"], self.bc["model"]
        title(ws, "Business case: route Self-Serve teams that look like Enterprise buyers to sales",
              "Three years, a new cohort of upgrades each year, upgrades mid-year. US$. Starting point: Vimeo's Q3 2025 10-Q.")
        widths(ws, {"A": 48, "B": 16, "C": 16, "D": 16, "E": 16})
        put(ws, "A4", "Starting point (Q3 2025)", bold=True, size=11)
        start = [("Self-Serve subscribers", base["self_serve_subscribers"], NUM, f"{self.lk(None, R_['sg_subs'], '2025', '2025Q3', 'Self-Serve')}*1000"),
                 ("Self-Serve ARPU", base["self_serve_arpu"], USD, self.lk(None, R_['sg_arpu'], '2025', '2025Q3', 'Self-Serve')),
                 ("Enterprise ARPU", base["enterprise_arpu"], USD, self.lk(None, R_['sg_arpu'], '2025', '2025Q3', 'Vimeo Enterprise'))]
        for r, (label, v, fmt, f) in enumerate(start, start=5):
            put(ws, f"A{r}", label)
            put(ws, f"B{r}", f"={f}", color=GREEN, fmt=fmt)
            self.check("Business case", label, v, f"'Business case'!B{r}")
        put(ws, "A9", "The bet", bold=True, size=11)
        derived = [("Eligible accounts", "=B5*" + R_["eligible_share"], "eligible", NUM),
                   ("Upgrades per year", "=B10*" + R_["uplift"], "upgrades", NUM1),
                   ("First-year contract value", "=B7*" + R_["entry_acv_share"], "acv", USD),
                   ("Account executives (hired to plan)", f"=ROUNDUP(B10*{R_['planned_uplift']}/{R_['upgrades_per_ae']}-1E-9,0)", None, NUM),
                   ("Self-Serve revenue lost per upgrade", f"=B6*{R_['cannibalised']}", None, USD)]
        for r, (label, f, key, fmt) in enumerate(derived, start=10):
            put(ws, f"A{r}", label)
            put(ws, f"B{r}", f, fmt=fmt, bold=True)
            if key:
                self.check("Business case", label, m[key], f"'Business case'!B{r}")
        header(ws, 16, ["US$", "Year 1", "Year 2", "Year 3", "Total"], height=20)
        y = m["years"]
        lines = ["Revenue weight (cohort-years)", "Enterprise revenue", "Self-Serve revenue lost", "Net new revenue",
                 "Gross profit", "Account executives", "Outreach", "Build", "Contribution", "Cumulative", "Discount factor",
                 "Present value"]
        for k, label in enumerate(lines):
            put(ws, f"A{17 + k}", label, bold=label in ("Net new revenue", "Contribution", "Present value"))
        ret = R_["retention"]
        for t in (1, 2, 3):
            c = col(1 + t)
            w = "+".join(["0.5" if cc == t else f"{ret}^{t - cc}" for cc in range(1, t + 1)])
            n_full = t - 1
            put(ws, f"{c}17", f"={w}", fmt=NUM2)
            put(ws, f"{c}18", f"=$B$11*$B$12*{c}17", fmt=USD)
            put(ws, f"{c}19", f"=$B$11*$B$14*({n_full}+0.5)", fmt=USD)
            put(ws, f"{c}20", f"={c}18-{c}19", fmt=USD, bold=True)
            put(ws, f"{c}21", f"={c}20*{R_['gross_margin']}", fmt=USD)
            put(ws, f"{c}22", f"=$B$13*{R_['ae_cost']}", fmt=USD)
            put(ws, f"{c}23", f"=$B$10*{R_['outreach_cost']}", fmt=USD)
            put(ws, f"{c}24", f"={R_['build_cost']}" if t == 1 else 0, fmt=USD)
            put(ws, f"{c}25", f"={c}21-{c}22-{c}23-{c}24", fmt=USD, bold=True)
            put(ws, f"{c}26", f"={c}25" if t == 1 else f"={col(t)}26+{c}25", fmt=USD)
            put(ws, f"{c}27", f"=1/(1+{R_['discount_rate']})^{t}", fmt="0.0000")
            put(ws, f"{c}28", f"={c}25*{c}27", fmt=USD, bold=True)
            row = y.iloc[t - 1]
            for rr, name in ((18, "enterprise_revenue"), (19, "cannibalised"), (20, "net_revenue"), (21, "gross_profit"),
                             (22, "sales_cost"), (23, "outreach_cost"), (25, "contribution"), (28, "present_value")):
                self.check("Business case", f"Year {t} {name}", float(row[name]), f"'Business case'!{c}{rr}")
        for rr in (18, 19, 20, 21, 22, 23, 24, 25, 28):
            put(ws, f"E{rr}", f"=SUM(B{rr}:D{rr})", fmt=USD, bold=rr in (20, 25, 28))
        put(ws, "A30", "Net present value", bold=True, size=11)
        put(ws, "B30", "=E28", fmt=USD, bold=True)
        put(ws, "A31", "Payback year (first year cumulative is positive)")
        put(ws, "B31", '=IF(B26>=0,1,IF(C26>=0,2,IF(D26>=0,3,"after year 3")))')
        put(ws, "A32", "Year-3 net new revenue")
        put(ws, "B32", "=D20", fmt=USD)
        # break-even: with the team hired to plan, net present value is linear in the uplift
        put(ws, "A34", "Break-even uplift", bold=True, size=11)
        put(ws, "A35", "Present value of gross profit per unit of uplift")
        put(ws, "B35", f"=SUMPRODUCT(B21:D21,B27:D27)/{R_['uplift']}", fmt=USD)
        put(ws, "A36", "Present value of fixed costs")
        put(ws, "B36", "=SUMPRODUCT(B22:D22,B27:D27)+SUMPRODUCT(B23:D23,B27:D27)+SUMPRODUCT(B24:D24,B27:D27)", fmt=USD)
        put(ws, "A37", "Uplift at which net present value is zero", bold=True)
        put(ws, "B37", "=B36/B35", fmt=PCT3, bold=True, fill=YELLOW)
        self.check("Business case", "Net present value", m["npv"], "'Business case'!B30")
        self.check("Business case", "Payback year", m["payback_year"], "'Business case'!B31")
        self.check("Business case", "Break-even uplift (closed form vs bisection)", self.bc["breakeven_uplift"], "'Business case'!B37")
        self.R.update({"npv": "'Business case'!$B$30", "payback": "'Business case'!$B$31", "breakeven": "'Business case'!$B$37",
                       "eligible": "'Business case'!$B$10", "year3": "'Business case'!$B$32"})

    def sensitivity(self, ws):
        R_, s = self.R, self.bc["sensitivity"]
        title(ws, "Sensitivity: net present value with one driver at its low or high value",
              "Each row recomputes the three-year case in full. Sorted by swing at the time of building; values update live.")
        drivers = list(C.SENSITIVITY)
        heads = ["Driver", "Case", *drivers, "Eligible", "Upgrades", "ACV", "AEs", "W1", "W2", "W3", "PV1", "PV2", "PV3", "NPV"]
        header(ws, 4, heads, height=32)
        widths(ws, {"A": 18, "B": 8, **{col(k): 11 for k in range(3, len(heads) + 1)}})
        r = 5
        dcol = {d: col(3 + i) for i, d in enumerate(drivers)}
        n = len(drivers)
        E, U, A, AE = col(3 + n), col(4 + n), col(5 + n), col(6 + n)
        W = [col(7 + n + i) for i in range(3)]
        PV = [col(10 + n + i) for i in range(3)]
        NPV = col(13 + n)
        self.sens_rows = {}
        for row in s.itertuples(index=False):
            for case, value in (("low", row.low_value), ("high", row.high_value)):
                put(ws, f"A{r}", row.driver, color=BLUE)
                put(ws, f"B{r}", case, color=BLUE)
                for d in drivers:
                    if d == row.driver:
                        put(ws, f"{dcol[d]}{r}", float(value), color=BLUE, fill=YELLOW)
                    else:
                        put(ws, f"{dcol[d]}{r}", f"={R_[d]}", color=GREEN)
                ret = f"{dcol['retention']}{r}"
                put(ws, f"{E}{r}", f"='Business case'!$B$5*{dcol['eligible_share']}{r}", fmt=NUM)
                put(ws, f"{U}{r}", f"={E}{r}*{dcol['uplift']}{r}", fmt=NUM1)
                put(ws, f"{A}{r}", f"='Business case'!$B$7*{dcol['entry_acv_share']}{r}", fmt=USD)
                put(ws, f"{AE}{r}", f"=ROUNDUP({E}{r}*{R_['planned_uplift']}/{dcol['upgrades_per_ae']}{r}-1E-9,0)", fmt=NUM)
                for t in (1, 2, 3):
                    put(ws, f"{W[t - 1]}{r}", "=" + "+".join(["0.5" if cc == t else f"{ret}^{t - cc}" for cc in range(1, t + 1)]), fmt=NUM2)
                    gp = f"({U}{r}*{A}{r}*{W[t - 1]}{r}-{U}{r}*'Business case'!$B$14*({t - 1}+0.5))*{dcol['gross_margin']}{r}"
                    cost = f"{AE}{r}*{R_['ae_cost']}+{E}{r}*{dcol['outreach_cost']}{r}" + (f"+{R_['build_cost']}" if t == 1 else "")
                    put(ws, f"{PV[t - 1]}{r}", f"=({gp}-({cost}))/(1+{R_['discount_rate']})^{t}", fmt=USD)
                put(ws, f"{NPV}{r}", f"=SUM({PV[0]}{r}:{PV[2]}{r})", fmt=USD, bold=True)
                self.check("Sensitivity", f"{row.driver} {case}", float(row.npv_low if case == "low" else row.npv_high), f"Sensitivity!{NPV}{r}")
                r += 1
        put(ws, f"A{r + 1}", "Base case", bold=True)
        put(ws, f"{NPV}{r + 1}", f"={R_['npv']}", color=GREEN, fmt=USD, bold=True)
        ws.conditional_formatting.add(f"{NPV}5:{NPV}{r - 1}", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
        ws.freeze_panes = "C5"

    # ---- The test ------------------------------------------------------------------------------
    def design(self, ws):
        R_, d = self.R, self.ex["designs"]
        title(ws, "Test design: can the test see the break-even uplift?",
              "Per-arm sample size for two proportions (normal approximation). The bar is the break-even uplift from the business case.")
        widths(ws, {"A": 42, "B": 10, "C": 10, "D": 13, "E": 13, "F": 13, "G": 12, "H": 12, "I": 12, "J": 12, "K": 10})
        header(ws, 4, ["Design", "Window (days)", "One-sided (1 = yes)", "Eligibility multiplier", "Control rate",
                       "Treatment rate at break-even", "z alpha", "z power", "Per arm", "Needed", "Available", "Feasible"])
        widths(ws, {"L": 10})
        for r, t in enumerate(d.itertuples(index=False), start=5):
            mult = 2.0 if "doubled" in t.design else 1.0
            put(ws, f"A{r}", t.design, color=BLUE)
            put(ws, f"B{r}", int(t.window_days), color=BLUE)
            put(ws, f"C{r}", int(t.one_sided), color=BLUE)
            put(ws, f"D{r}", mult, color=BLUE)
            put(ws, f"E{r}", f"={R_['baseline_rate']}*B{r}/90", fmt=PCT3)
            put(ws, f"F{r}", f"=E{r}+{R_['breakeven']}/D{r}*B{r}/365", fmt=PCT3)
            put(ws, f"G{r}", f"=NORMSINV(1-{R_['alpha']}/IF(C{r}=1,1,2))", fmt="0.000")
            put(ws, f"H{r}", f"=NORMSINV({R_['power']})", fmt="0.000")
            put(ws, f"I{r}", f"=ROUNDUP((G{r}+H{r})^2*(E{r}*(1-E{r})+F{r}*(1-F{r}))/(F{r}-E{r})^2,0)", fmt=NUM, bold=True)
            put(ws, f"J{r}", f"=2*I{r}", fmt=NUM)
            put(ws, f"K{r}", f"={R_['eligible']}*D{r}", fmt=NUM)
            put(ws, f"L{r}", f'=IF(J{r}<=K{r},"Yes","No")', bold=True)
            self.check("Test design", f"{t.design} per arm", int(t.per_arm), f"'Test design'!I{r}")
            self.check("Test design", f"{t.design} feasible", "Yes" if t.feasible else "No", f"'Test design'!L{r}")
        last = 4 + len(d)
        ws.conditional_formatting.add(f"L5:L{last}", CellIsRule(operator="equal", formula=['"Yes"'], fill=GREEN_FILL))
        ws.conditional_formatting.add(f"L5:L{last}", CellIsRule(operator="equal", formula=['"No"'], fill=RED_FILL))
        put(ws, f"A{last + 2}", "Doubling eligibility brings in weaker accounts: the uplift per account is assumed to halve.",
            color=MUTED, italic=True)

    def readout(self, ws):
        R_, s = self.R, self.ex["simulation"]
        title(ws, "Test readout: paste the counts, read the decision",
              "Filled with one simulated run of the 180-day design (true uplift 1% a year). Simulated, not Vimeo data.")
        widths(ws, {"A": 46, "B": 16, "C": 16})
        header(ws, 4, ["", "Control", "Treatment"], height=20)
        rows = [("Accounts", "accounts"), ("Upgrades in the window", "upgrades"), ("Self-Serve cancellations in the window", "cancellations")]
        for r, (label, key) in enumerate(rows, start=5):
            put(ws, f"A{r}", label)
            put(ws, f"B{r}", s["control"][key], color=BLUE, fmt=NUM)
            put(ws, f"C{r}", s["treatment"][key], color=BLUE, fmt=NUM)
        put(ws, "A8", "Window, days")
        put(ws, "B8", int(s["design"]["window_days"]), color=BLUE)
        calc = [
            ("Upgrade rate, control", "=B6/B5", PCT3, ("control_rate", s["control"]["rate"])),
            ("Upgrade rate, treatment", "=C6/C5", PCT3, ("treatment_rate", s["treatment"]["rate"])),
            ("Difference", "=B11-B10", PCT3, ("difference", s["difference"])),
            ("Standard error", "=SQRT(B10*(1-B10)/B5+B11*(1-B11)/C5)", "0.00000", ("se", s["se"])),
            ("One-sided 95% lower bound", f"=B12-NORMSINV(1-{R_['alpha']})*B13", PCT3, ("lower_bound", s["lower_bound"])),
            ("Break-even difference in the window", f"={R_['breakeven']}*B8/365", PCT3, ("breakeven_window", s["breakeven_window"])),
            ("Uplift per year, estimate", "=B12*365/B8", PCT2, ("uplift_yearly_estimate", s["uplift_yearly_estimate"])),
            ("Sample ratio check, p-value", "=2*(1-NORMSDIST(SQRT(((B5-(B5+C5)/2)^2+(C5-(B5+C5)/2)^2)/((B5+C5)/2))))", "0.000", ("srm_p", s["srm_p"])),
            ("Cancellations, difference", "=C7/C5-B7/B5", PCT2, ("guardrail_difference", s["guardrail_difference"])),
            ("Cancellations, upper bound", f"=B18+NORMSINV(1-{R_['alpha']})*SQRT(B7/B5*(1-B7/B5)/B5+C7/C5*(1-C7/C5)/C5)", PCT2,
             ("guardrail_upper", s["guardrail_upper"])),
            ("Largest rise accepted in the window", f"={R_['guardrail_margin']}*B8/90", PCT2, None),
        ]
        for r, (label, f, fmt, chk) in enumerate(calc, start=10):
            put(ws, f"A{r}", label)
            put(ws, f"B{r}", f, fmt=fmt)
            if chk:
                self.check("Test readout", label, float(chk[1]), f"'Test readout'!B{r}")
        put(ws, "A22", "Decision", bold=True, size=12)
        put(ws, "B22", ('=IF(B17<0.001,"Stop: the split is broken, fix assignment and rerun",'
                        'IF(B14<=0,"Stop or redesign: no detectable uplift",'
                        'IF(B19>=B20,"Do not scale: outreach may be pushing Self-Serve customers away",'
                        'IF(B12>=B15,"Scale: the uplift is real and above break-even",'
                        '"Iterate: the uplift is real but below break-even; cheaper outreach or better targeting first"))))'),
            bold=True, size=12)
        self.check("Test readout", "Decision", s["decision"], "'Test readout'!B22")
        put(ws, "A24", "Rule, in order: a broken split invalidates the test; without an uplift there is nothing to protect; "
                       "then the guardrail; then the business bar.", color=MUTED, italic=True)

    def voi(self, ws):
        R_, v = self.R, self.ex["voi"]
        priors = list(C.PRIORS)
        title(ws, "Is the test worth running?",
              "Each option's expected value under a prior belief about the real uplift. NPV is linear in the uplift, so it comes "
              "straight from the business case; the chance that the test says 'Scale' comes from simulated tests (blue).")
        widths(ws, {"A": 34, "B": 13, "C": 13, "D": 16, "E": 14, "F": 16, "G": 16, "H": 14})
        put(ws, "A4", "Test window, days")
        put(ws, "B4", "=INDEX('Test design'!B5:B8,MATCH(\"Yes\",'Test design'!L5:L8,0))", color=GREEN)
        put(ws, "A5", "Cost of the test (build, outreach to the treated half, one account executive)")
        put(ws, "B5", f"={R_['eligible']}/2*{R_['outreach_cost']}*B4/365+{C.TEST_STAFF}*{R_['ae_cost']}*B4/365+{R_['build_cost']}", fmt=USD)
        put(ws, "A6", "Delay factor for value after the test")
        put(ws, "B6", f"=1/(1+{R_['discount_rate']})^(B4/365)", fmt="0.0000")
        self.check("Value of test", "Cost of the test", v[priors[0]]["test_cost"], "'Value of test'!B5")
        header(ws, 8, ["True uplift", *[f"Prior: {p}" for p in priors], "NPV if launched now", "Chance the test says Scale",
                       "Value if scaled after the test", "Test-first value", "NPV, if positive"], height=44)
        t = v[priors[0]]["table"]
        first = 9
        for r, row in enumerate(t.itertuples(index=False), start=first):
            put(ws, f"A{r}", float(row.uplift), color=BLUE, fmt=PCT2)
            for j, pr in enumerate(priors):
                put(ws, f"{col(2 + j)}{r}", float(C.PRIORS[pr][row.uplift]), color=BLUE, fmt=PCT)
            c = col(2 + len(priors))
            put(ws, f"{c}{r}", f"='Business case'!$B$35*A{r}-'Business case'!$B$36", fmt=USD)
            put(ws, f"{col(3 + len(priors))}{r}", float(row.p_scale), color=BLUE, fmt=PCT1)
            put(ws, f"{col(4 + len(priors))}{r}", f"=({c}{r}+{R_['build_cost']}/(1+{R_['discount_rate']}))*$B$6", fmt=USD)
            put(ws, f"{col(5 + len(priors))}{r}", f"={col(3 + len(priors))}{r}*{col(4 + len(priors))}{r}-$B$5", fmt=USD)
            put(ws, f"{col(6 + len(priors))}{r}", f"=MAX(0,{c}{r})", fmt=USD)
            self.check("Value of test", f"NPV at {row.uplift:.3f}", float(row.npv_launch_now), f"'Value of test'!{c}{r}")
            self.check("Value of test", f"Test-first value at {row.uplift:.3f}", float(row.test_value), f"'Value of test'!{col(5 + len(priors))}{r}")
        last = first + len(t) - 1
        rng = lambda letter: f"{letter}{first}:{letter}{last}"
        npv_c, tv_c, pos_c = col(2 + len(priors)), col(5 + len(priors)), col(6 + len(priors))
        r = last + 2
        header(ws, r, ["Under the prior", *[p for p in priors]], height=20)
        lines = [("Launch now", lambda pc: f"=SUMPRODUCT({rng(pc)},{rng(npv_c)})", "launch_now"),
                 ("Do nothing", lambda pc: "=0", "do_nothing"),
                 ("Test first, then follow the rule", lambda pc: f"=SUMPRODUCT({rng(pc)},{rng(tv_c)})", "test_first"),
                 ("Value of perfect information", lambda pc: f"=SUMPRODUCT({rng(pc)},{rng(pos_c)})-MAX(0,SUMPRODUCT({rng(pc)},{rng(npv_c)}))", "evpi")]
        for k, (label, f, key) in enumerate(lines, start=1):
            put(ws, f"A{r + k}", label, bold=key == "test_first")
            for j, pr in enumerate(priors):
                cc = col(2 + j)
                put(ws, f"{cc}{r + k}", f(cc), fmt=USD, bold=key == "test_first")
                self.check("Value of test", f"{pr} {label}", float(v[pr][key]), f"'Value of test'!{cc}{r + k}")
        k = len(lines) + 1
        put(ws, f"A{r + k}", "Value of running the test", bold=True)
        put(ws, f"A{r + k + 1}", "Best option", bold=True)
        for j, pr in enumerate(priors):
            cc = col(2 + j)
            put(ws, f"{cc}{r + k}", f"={cc}{r + 3}-MAX(0,{cc}{r + 1})", fmt=USD, bold=True, fill=YELLOW)
            put(ws, f"{cc}{r + k + 1}", f'=IF(AND({cc}{r + 3}>={cc}{r + 1},{cc}{r + 3}>=0),"Test first",IF({cc}{r + 1}>=0,"Launch now","Do nothing"))', bold=True)
            self.check("Value of test", f"{pr} value of running the test", float(v[pr]["value_of_test"]), f"'Value of test'!{cc}{r + k}")
            self.check("Value of test", f"{pr} best option", v[pr]["best"], f"'Value of test'!{cc}{r + k + 1}")
        put(ws, f"A{r + k + 3}", "The test's own upgrades are not counted, which understates its value slightly. The chance of 'Scale' at an "
                                 "uplift below break-even is the rule's cost: it can scale a bet that loses.", color=MUTED, italic=True)

    def summary(self, ws):
        R_ = self.R
        title(ws, "Summary", "Everything on this sheet is a link or a formula.")
        widths(ws, {"A": 52, "B": 20})
        rows = [("Eligible Self-Serve accounts", f"={R_['eligible']}", NUM),
                ("Net present value, 3 years", f"={R_['npv']}", USD),
                ("Payback year", f"={R_['payback']}", None),
                ("Year-3 net new revenue", f"={R_['year3']}", USD),
                ("Break-even uplift per eligible account per year", f"={R_['breakeven']}", PCT3),
                ("Planned uplift", f"={R_['uplift']}", PCT2),
                ("Smallest feasible test", "=INDEX('Test design'!A5:A8,MATCH(\"Yes\",'Test design'!L5:L8,0))", None),
                ("Decision on the simulated run", "='Test readout'!B22", None)]
        for r, (label, f, fmt) in enumerate(rows, start=4):
            put(ws, f"A{r}", label, bold=True)
            put(ws, f"B{r}", f, color=GREEN, fmt=fmt)
        self.check("Summary", "Smallest feasible test", experiment.chosen_design()["design"], "Summary!B10")

    def reconciliation(self, ws):
        title(ws, "Reconciliation: workbook formulas against the Python pipeline",
              "Column D was written by src/build_workbook.py from the pandas results; column E is the live formula.")
        widths(ws, {"A": 16, "B": 58, "C": 3, "D": 20, "E": 20, "F": 14, "G": 8})
        header(ws, 6, ["Area", "Item", "", "Python value", "Workbook value", "Difference", "Match"])
        first, last = 7, 6 + len(self.checks)
        put(ws, "B3", (f'=IF(COUNTIF(G{first}:G{last},"No")=0,"All "&COUNTA(B{first}:B{last})&" checks match",'
                       f'COUNTIF(G{first}:G{last},"No")&" of "&COUNTA(B{first}:B{last})&" checks do not match")'), bold=True, size=12)
        for r, (area, item, value, ref) in enumerate(self.checks, start=first):
            put(ws, f"A{r}", area)
            put(ws, f"B{r}", item)
            fmt = "0.000000" if isinstance(value, float) else None
            put(ws, f"D{r}", value, color=BLUE, fmt=fmt)
            put(ws, f"E{r}", f"={ref}", color=GREEN, fmt=fmt)
            put(ws, f"F{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")', fmt="0.0E+00")
            put(ws, f"G{r}", (f'=IF(AND(D{r}="",E{r}=""),"Yes",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),'
                              f'IF(ABS(E{r}-D{r})<=1E-6*MAX(1,ABS(D{r})),"Yes","No"),IF(D{r}=E{r},"Yes","No")))'))
        for text, fill in (("Yes", GREEN_FILL), ("No", RED_FILL)):
            ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill))
        ws.freeze_panes = "A7"
        self.recon = "Reconciliation!B3"

    def cover(self, ws, names):
        widths(ws, {"A": 3, "B": 22, "C": 100})
        put(ws, "B2", "Vimeo: the Self-Serve to Enterprise upgrade case", bold=True, size=18)
        put(ws, "B3", "Public filings, a business case and the test that would decide it", color=MUTED, size=12)
        put(ws, "B5", "How to read it", bold=True, size=11)
        for i, (label, meaning, colour, fill) in enumerate([
                ("Blue text", "an input: from the filings, or an assumption", BLUE, None), ("Black text", "a formula", INK, None),
                ("Green text", "a link to another sheet", GREEN, None), ("Yellow fill", "an assumption the result depends on", INK, YELLOW)], start=6):
            put(ws, f"B{i}", label, color=colour, fill=fill)
            put(ws, f"C{i}", meaning)
        put(ws, "B11", "Sheets", bold=True, size=11)
        purpose = {"Summary": "The answer in eight lines.", "Business case": "Three years of the bet, its value and its break-even uplift.",
                   "Sensitivity": "The case recomputed with each driver at its low and high value.",
                   "Test design": "Sample sizes for four designs against the accounts available.",
                   "Test readout": "Paste the counts of a real test and read the decision.",
                   "Value of test": "Launch now, do nothing, or test first: each option's value under two priors.",
                   "Bridge": "Each category's revenue change split into volume and price.",
                   "Segments": "Subscribers, ARPU, bookings and revenue from the filings.",
                   "P&L": "Revenue, costs and margins from the XBRL facts.", "Inputs": "Every assumption, with its reason.",
                   "Reconciliation": "Each figure checked against the Python pipeline."}
        for i, n in enumerate(names[1:], start=12):
            c = put(ws, f"B{i}", n, color="1F5FA8")
            c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{n}'!A1")
            put(ws, f"C{i}", purpose[n])
        r = 12 + len(names)
        put(ws, f"B{r}", "Check", bold=True, size=11)
        put(ws, f"C{r}", f"={self.recon}", color=GREEN, bold=True)
        put(ws, f"B{r + 2}", "Data", bold=True, size=11)
        put(ws, f"C{r + 2}", "Vimeo, Inc. 10-K and 10-Q filings with the SEC, 2021 to Q3 2025. Assumptions are mine and labelled. "
                             "The test readout is simulated. Nothing here uses non-public information.", wrap=True)
        ws.row_dimensions[r + 2].height = 28
        put(ws, f"B{r + 4}", "Domenico Perroni", color=MUTED)
        c = put(ws, f"C{r + 4}", REPO, color="1F5FA8")
        c.hyperlink = REPO


def build(a=None, bc=None, ex=None, path: Path = OUTPUT) -> Workbook_:
    a = analysis.run() if a is None else a
    bc = business_case.run() if bc is None else bc
    ex = experiment.run() if ex is None else ex
    w = Workbook_(a, bc, ex)
    w.build(path)
    return w


def main() -> None:
    w = build()
    print(f"workbook: {len(w.checks)} reconciliation checks -> {OUTPUT.relative_to(C.ROOT)}")


if __name__ == "__main__":
    main()
